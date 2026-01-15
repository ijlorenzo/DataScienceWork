import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from datetime import datetime

class ExcelGenerator:
    """Generate formatted Excel output with LLM reasoning fields"""

    OUTPUT_FIELDS = [
        # DESCRIBE
        'Customer need', 'Customer need - LLM Reasoning',
        'Product/feature request',
        'Need description',
        'Request type', 'Request type - LLM Reasoning',

        # PRODUCT
        'Product Area', 'Product Area - LLM Reasoning',
        'Product Category', 'Product Category - LLM Reasoning',
        'Customer-facing product category', 'Customer-facing product category - LLM Reasoning',

        # CUSTOMER IMPACT
        'Vertical', 'Vertical - LLM Reasoning',
        '# of customers', '# of customers - LLM Reasoning',
        'Personas', 'Personas - LLM Reasoning',
        'UXR Needs Based Segmentation', 'UXR Needs Based Segmentation - LLM Reasoning',
        'Sales Segment', 'Sales Segment - LLM Reasoning',
        'User Touchpoint', 'User Touchpoint - LLM Reasoning',

        # PRIORITY
        'Rev Blocker [Flag]',
        'Revenue [Tshirt Size]', 'Revenue [Tshirt Size] - LLM Reasoning',
        'Reach [T-Shirt Size]', 'Reach [T-Shirt Size] - LLM Reasoning',
        'Revenue Sizing Needed',
        'Revenue Estimate [Where Avail]',
        'Severity [T-Shirt Size]', 'Severity [T-Shirt Size] - LLM Reasoning',

        # TRACKING
        'Roadmap Status', 'Last Updated (mm/dd/yy)', 'ADO Link', 'Unique ID',

        # SOURCE
        'Listening Type', 'Listening Type - LLM Reasoning',
        'Listening Source', 'Date Created (mm/dd/yy)',
        'CX Insights Priority Needs (y/n)', 'Growth Blocker List (y/n)',

        # SUPPORTING MATERIALS
        'One pager', 'BRD', 'Transcripts', 'Recap', 'Ideas portal link',

        # OWNERSHIP
        'PMM', 'PM', 'GPM', 'Business Planning'
    ]

    def generate(self, items, output_path):
        """Generate Excel file with formatting"""
        # Create DataFrame
        df = pd.DataFrame(items)

        # Reorder columns
        available_cols = [col for col in self.OUTPUT_FIELDS if col in df.columns]
        df = df[available_cols]

        # Write to Excel
        df.to_excel(output_path, index=False, sheet_name='Customer Needs')

        # Apply formatting
        wb = load_workbook(output_path)
        ws = wb['Customer Needs']

        # Freeze header row
        ws.freeze_panes = 'A2'

        # Format header
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        # Color LLM reasoning columns
        reasoning_fill = PatternFill(start_color='E7F3FF', end_color='E7F3FF', fill_type='solid')
        for idx, col in enumerate(available_cols, 1):
            if 'LLM Reasoning' in col:
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, idx).fill = reasoning_fill

        # Color-code severity
        severity_col = available_cols.index('Severity [T-Shirt Size]') + 1 if 'Severity [T-Shirt Size]' in available_cols else None
        if severity_col:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, severity_col)
                if cell.value == 'High':
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                elif cell.value == 'Medium':
                    cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                elif cell.value == 'Low':
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

        # Auto-fit columns
        for column in ws.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column[0].column_letter].width = adjusted_width

        wb.save(output_path)
        return output_path
